
import argparse
import os

import matplotlib
matplotlib.use("Agg")                                   # guarda archivos sin abrir ventanas
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.colors import LinearSegmentedColormap
from matplotlib.ticker import FuncFormatter, MaxNLocator
from openpyxl import load_workbook

# ===== Estilo =====
COLOR = {"CONV1D": "#2a78d6", "CONV2D": "#eb6834"}     # azul / naranja (seguros para daltonismo)
TEXTO, TEXTO_2, GRILLA = "#0b0b0b", "#52514e", "#e4e3df"
AZULES = LinearSegmentedColormap.from_list(            # escala secuencial de un solo tono
    "azules", ["#fcfcfb", "#cde2fb", "#86b6ef", "#3987e5", "#256abf", "#184f95", "#0d366b"])
plt.rcParams.update({
    "font.family": "DejaVu Sans", "font.size": 10, "axes.titlesize": 12, "axes.titleweight": "bold",
    "axes.edgecolor": GRILLA, "axes.labelcolor": TEXTO_2, "xtick.color": TEXTO_2, "ytick.color": TEXTO_2,
    "axes.spines.top": False, "axes.spines.right": False, "axes.grid": True, "grid.color": GRILLA,
    "grid.linewidth": 0.8, "axes.axisbelow": True, "legend.frameon": False, "figure.dpi": 100,
})


def leer_hoja(wb, nombre):
    """Hoja de Excel -> (encabezados, filas)."""
    filas = [list(f) for f in wb[nombre].iter_rows(values_only=True) if any(v is not None for v in f)]
    return filas[0], filas[1:]


def como_dicts(wb, nombre):
    enc, filas = leer_hoja(wb, nombre)
    return [dict(zip(enc, f)) for f in filas]


def guardar(fig, carpeta, archivo):
    ruta = os.path.join(carpeta, archivo)
    fig.savefig(ruta, dpi=200, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    print("  ", ruta)


def etiquetar_barras(ax, barras, fmt="{:.1f}", dy=1.0):
    for b in barras:
        h = b.get_height()
        ax.text(b.get_x() + b.get_width() / 2, h + dy, fmt.format(h), ha="center", va="bottom",
                fontsize=9, color=TEXTO)


# ===== Gráficos =====
def g_accuracy_f1(wb, carpeta):
    res = como_dicts(wb, "Resumen")
    metricas = [("Accuracy", "Accuracy media (%)", "Accuracy desv. (%)"),
                ("F1 macro", "F1 macro media (%)", "F1 macro desv. (%)")]
    x = np.arange(len(metricas))
    ancho = 0.36
    fig, ax = plt.subplots(figsize=(6.5, 4.2))
    for k, r in enumerate(res):
        medias = [r[m] for _, m, _ in metricas]
        desv = [r[d] for _, _, d in metricas]
        b = ax.bar(x + (k - 0.5) * (ancho + 0.02), medias, ancho, yerr=desv, capsize=4,
                   color=COLOR[r["Modelo"]], label=r["Modelo"].replace("CONV", "Conv"),
                   error_kw={"elinewidth": 1, "ecolor": TEXTO_2})
        etiquetar_barras(ax, b, dy=max(desv) + 1)
    ax.set_xticks(x, [m for m, _, _ in metricas])
    ax.set_ylim(0, 100)
    ax.set_ylabel("%")
    ax.grid(axis="x", visible=False)
    ax.set_title(f"Rendimiento en el conjunto de prueba ({res[0]['N° semillas']} particiones)", loc="left")
    ax.legend(loc="upper right", ncols=2)
    guardar(fig, carpeta, "01_accuracy_f1.png")


def g_por_semilla(wb, carpeta):
    filas = como_dicts(wb, "Por semilla")
    semillas = sorted({r["Semilla"] for r in filas})
    x = np.arange(len(semillas))
    ancho = 0.36
    fig, ax = plt.subplots(figsize=(7, 4))
    for k, m in enumerate(("CONV1D", "CONV2D")):
        val = [next(r["Accuracy (%)"] for r in filas if r["Modelo"] == m and r["Semilla"] == s) for s in semillas]
        b = ax.bar(x + (k - 0.5) * (ancho + 0.02), val, ancho, color=COLOR[m], label=m.replace("CONV", "Conv"))
        etiquetar_barras(ax, b)
    ax.set_xticks(x, [f"Semilla {s}" for s in semillas])
    ax.set_ylim(0, 100)
    ax.set_ylabel("Accuracy (%)")
    ax.grid(axis="x", visible=False)
    ax.set_title("Accuracy por partición (misma partición para ambos modelos)", loc="left")
    ax.legend(loc="upper right", ncols=2)
    guardar(fig, carpeta, "02_accuracy_por_semilla.png")


def g_f1_por_sena(wb, carpeta):
    filas = sorted(como_dicts(wb, "Por seña"), key=lambda r: r["F1 Conv2D (%)"])
    senas = [r["Seña"] for r in filas]
    y = np.arange(len(senas))
    alto = 0.38
    fig, ax = plt.subplots(figsize=(7, 0.5 * len(senas) + 1.2))
    for k, (m, col) in enumerate((("CONV1D", "F1 Conv1D (%)"), ("CONV2D", "F1 Conv2D (%)"))):
        val = [r[col] for r in filas]
        b = ax.barh(y + (0.5 - k) * (alto + 0.02), val, alto, color=COLOR[m], label=m.replace("CONV", "Conv"))
        for bb, v in zip(b, val):
            ax.text(v + 1, bb.get_y() + bb.get_height() / 2, f"{v:.0f}", va="center", fontsize=8, color=TEXTO_2)
    ax.set_yticks(y, senas)
    ax.set_xlim(0, 105)
    ax.set_xlabel("F1 (%)")
    ax.grid(axis="y", visible=False)
    ax.set_title("F1 por seña (promedio de las particiones)", loc="left")
    ax.legend(loc="lower right", ncols=2)
    guardar(fig, carpeta, "03_f1_por_sena.png")


def g_matriz(wb, carpeta, modelo, archivo):
    enc, filas = leer_hoja(wb, f"Matriz {modelo.replace('CONV', 'Conv')}")
    clases = enc[1:enc.index("Total")]
    cm = np.array([[f[1 + j] for j in range(len(clases))] for f in filas], dtype=float)
    pct = 100 * cm / np.maximum(cm.sum(1, keepdims=True), 1)          # % de cada fila (seña real)
    fig, ax = plt.subplots(figsize=(7.2, 6.2))
    im = ax.imshow(pct, cmap=AZULES, vmin=0, vmax=100)
    for i in range(len(clases)):
        for j in range(len(clases)):
            if cm[i, j] > 0:
                ax.text(j, i, f"{pct[i, j]:.0f}", ha="center", va="center", fontsize=8,
                        color="white" if pct[i, j] > 55 else TEXTO, fontweight="bold" if i == j else "normal")
    ax.set_xticks(range(len(clases)), clases, rotation=45, ha="right")
    ax.set_yticks(range(len(clases)), clases)
    ax.set_xlabel("Seña predicha")
    ax.set_ylabel("Seña real")
    ax.grid(False)
    for s in ax.spines.values():
        s.set_visible(False)
    cb = fig.colorbar(im, ax=ax, fraction=0.046, pad=0.03)
    cb.set_label("% de la seña real")
    cb.outline.set_visible(False)
    acc = 100 * np.trace(cm) / cm.sum()
    ax.set_title(f"Matriz de confusión {modelo.replace('CONV', 'Conv')} — accuracy global {acc:.1f}%\n"
                 f"(suma de todas las particiones, {int(cm.sum())} predicciones)", loc="left")
    guardar(fig, carpeta, archivo)


def g_curvas(wb, carpeta, modelo, prefijo):
    """Un gráfico por semilla con las curvas de aprendizaje: loss (izquierda) y accuracy (derecha)."""
    todas = [r for r in como_dicts(wb, "Historial") if r["Modelo"] == modelo]
    nombre = modelo.replace("CONV", "Conv")
    for s in sorted({r["Semilla"] for r in todas}):
        filas = [r for r in todas if r["Semilla"] == s]
        ep = [r["Época"] for r in filas]
        fig, (a1, a2) = plt.subplots(1, 2, figsize=(10, 3.8))
        for ax, clave, sufijo, titulo in ((a1, "Loss", "", "Pérdida (loss)"),
                                          (a2, "Accuracy", " (%)", "Accuracy (%)")):
            ent = [r[f"{clave} entrenamiento{sufijo}"] for r in filas]
            val = [r[f"{clave} validación{sufijo}"] for r in filas]
            ax.plot(ep, ent, color="#2a78d6", lw=2, label="Entrenamiento")
            ax.plot(ep, val, color="#eb6834", lw=2, label="Validación")
            ax.set_xlabel("Época")
            ax.xaxis.set_major_locator(MaxNLocator(integer=True))
            ax.set_title(titulo, loc="left")
            if clave == "Loss":                          # recorta picos aislados (> 5 veces la mediana)
                med = float(np.median(ent + val))
                tope = max(v for v in ent + val if v <= 5 * med) * 1.08
                ax.set_ylim(0, tope)
                for e, v in zip(ep, val):
                    if v > tope:
                        ax.annotate(f"pico {v:.1f}", (e, tope), xytext=(6, -25), textcoords="offset points",
                                    fontsize=8, color=TEXTO_2, arrowprops={"arrowstyle": "-", "color": TEXTO_2})
        a2.set_ylim(0, 100)
        fig.suptitle(f"Curvas de aprendizaje {nombre} — semilla {s}", x=0.01, ha="left", fontweight="bold")
        fig.tight_layout(rect=(0, 0, 1, 0.97))
        fig.legend(*a1.get_legend_handles_labels(), loc="upper right", ncols=2, bbox_to_anchor=(1, 1.0))
        guardar(fig, carpeta, f"{prefijo}_curvas_{modelo.lower()}_semilla{s}.png")

def g_eficiencia(wb, carpeta):
    res = como_dicts(wb, "Resumen")
    paneles = [("Parámetros", "Parámetros", "{:,.0f}"), ("Tamaño TFLite (KB)", "TFLite (KB)", "{:,.0f}"),
               ("Inferencia (ms por seña, CPU)", "Inferencia (ms/seña)", "{:.2f}")]
    fig, axes = plt.subplots(1, 3, figsize=(10, 3.4))
    for ax, (titulo, col, fmt) in zip(axes, paneles):
        modelos = [r["Modelo"] for r in res]
        val = [r[col] for r in res]
        b = ax.bar([m.replace("CONV", "Conv") for m in modelos], val, 0.55, color=[COLOR[m] for m in modelos])
        for bb, v in zip(b, val):
            ax.text(bb.get_x() + bb.get_width() / 2, v * 1.02, fmt.format(v), ha="center", va="bottom",
                    fontsize=9, color=TEXTO)
        ax.set_ylim(0, max(val) * 1.18)
        ax.set_title(titulo, loc="left", fontsize=10)
        ax.grid(axis="x", visible=False)
        ax.tick_params(axis="y", labelsize=8)
        if col == "Parámetros":
            ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:,.0f}"))
    fig.suptitle("Costo computacional de cada modelo", x=0.01, ha="left", fontweight="bold")
    fig.tight_layout()
    guardar(fig, carpeta, "08_eficiencia.png")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--excel", default="data/resultados.xlsx")
    ap.add_argument("--salida", default="data/graficos")
    args = ap.parse_args()

    os.makedirs(args.salida, exist_ok=True)
    wb = load_workbook(args.excel, read_only=True)
    print(f"Leyendo {args.excel}\nGuardando en {args.salida}/:")
    g_accuracy_f1(wb, args.salida)
    g_por_semilla(wb, args.salida)
    g_f1_por_sena(wb, args.salida)
    g_matriz(wb, args.salida, "CONV1D", "04_matriz_conv1d.png")
    g_matriz(wb, args.salida, "CONV2D", "05_matriz_conv2d.png")
    g_curvas(wb, args.salida, "CONV1D", "06")
    g_curvas(wb, args.salida, "CONV2D", "07")
    g_eficiencia(wb, args.salida)
    print("Listo.")


if __name__ == "__main__":
    main()