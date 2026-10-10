
import csv
from collections import defaultdict

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

MODELOS = ("conv1d", "conv2d")
SALIDA = "data/resultados.xlsx"


def leer(ruta):
    with open(ruta, encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def pct(v):
    return round(float(v) * 100, 1)


res = {m: leer(f"data/resultados_{m}.csv") for m in MODELOS}
sen = {m: leer(f"data/por_sena_{m}.csv") for m in MODELOS}

semillas = [r["semilla"] for r in res["conv1d"]]
if semillas != [r["semilla"] for r in res["conv2d"]]:
    print("OJO: los dos modelos se corrieron con semillas distintas -> la comparación no es justa")

# ---- Hoja 1: Resumen
resumen = [["Modelo", "Accuracy media (%)", "Accuracy desv. (%)", "F1 macro media (%)", "F1 macro desv. (%)",
            "Parámetros", "TFLite (KB)", "Inferencia (ms/seña)", "Épocas (media)", "N° semillas"]]
for m in MODELOS:
    r = res[m]
    a = np.array([pct(x["accuracy"]) for x in r])
    f = np.array([pct(x["f1_macro"]) for x in r])
    resumen.append([m.upper(), round(a.mean(), 1), round(a.std(), 1), round(f.mean(), 1), round(f.std(), 1),
                    int(r[0]["parametros"]), round(float(r[0]["tflite_kb"])),
                    round(np.mean([float(x["ms_por_sena"]) for x in r]), 2),
                    round(np.mean([int(x["epocas"]) for x in r])), len(r)])

# ---- Hoja 2: Por semilla
por_semilla = [["Modelo", "Semilla", "Accuracy (%)", "F1 macro (%)", "Épocas", "Minutos"]]
for m in MODELOS:
    for x in res[m]:
        por_semilla.append([m.upper(), int(x["semilla"]), pct(x["accuracy"]), pct(x["f1_macro"]),
                            int(x["epocas"]), float(x["minutos"])])

# ---- Hoja 3: Por seña (promedio de las semillas)
prom = defaultdict(lambda: defaultdict(list))          # prom[seña][(modelo, métrica)] = [valores]
for m in MODELOS:
    for x in sen[m]:
        for met in ("precision", "recall", "f1"):
            prom[x["sena"]][(m, met)].append(float(x[met]))
por_sena = [["Seña", "Precisión Conv1D (%)", "Recall Conv1D (%)", "F1 Conv1D (%)",
             "Precisión Conv2D (%)", "Recall Conv2D (%)", "F1 Conv2D (%)", "Diferencia F1 (2D - 1D)"]]
for s in sorted(prom):
    v = {k: pct(np.mean(lst)) for k, lst in prom[s].items()}
    por_sena.append([s, v[("conv1d", "precision")], v[("conv1d", "recall")], v[("conv1d", "f1")],
                     v[("conv2d", "precision")], v[("conv2d", "recall")], v[("conv2d", "f1")],
                     round(v[("conv2d", "f1")] - v[("conv1d", "f1")], 1)])

# ---- Hoja 4: Detalle (crudo)
detalle = [["Modelo", "Semilla", "Seña", "Precisión (%)", "Recall (%)", "F1 (%)", "N° muestras de prueba"]]
for m in MODELOS:
    for x in sen[m]:
        detalle.append([m.upper(), int(x["semilla"]), x["sena"], pct(x["precision"]), pct(x["recall"]),
                        pct(x["f1"]), int(x["soporte"])])

# ---- Hojas 5 y 6: matrices de confusión (suma de las semillas)
matrices = {}
for m in MODELOS:
    with open(f"data/matriz_{m}.csv", encoding="utf-8") as fh:
        filas_m = list(csv.reader(fh))
    tabla = [["Real / Predicho"] + filas_m[0][1:] + ["Total", "Aciertos (%)"]]
    for f in filas_m[1:]:
        v = [int(x) for x in f[1:]]
        i = filas_m[0][1:].index(f[0])
        tabla.append([f[0]] + v + [sum(v), round(100 * v[i] / max(1, sum(v)), 1)])
    matrices[m] = tabla

# ---- Hoja 7: historial de entrenamiento
historial = [["Modelo", "Semilla", "Época", "Loss entrenamiento", "Accuracy entrenamiento (%)",
              "Loss validación", "Accuracy validación (%)"]]
for m in MODELOS:
    for x in leer(f"data/historial_{m}.csv"):
        historial.append([m.upper(), int(x["semilla"]), int(x["epoca"]), float(x["loss"]), pct(x["accuracy"]),
                          float(x["val_loss"]), pct(x["val_accuracy"])])

# ---- Escribir el Excel
wb = Workbook()
wb.remove(wb.active)
for nombre, tabla in (("Resumen", resumen), ("Por semilla", por_semilla),
                      ("Por seña", por_sena), ("Detalle señas", detalle),
                      ("Matriz Conv1D", matrices["conv1d"]), ("Matriz Conv2D", matrices["conv2d"]),
                      ("Historial", historial)):
    ws = wb.create_sheet(nombre)
    for fila in tabla:
        ws.append(fila)
    for c in ws[1]:                                      # encabezado en negrita con fondo
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for col in ws.columns:                               # ancho de columnas
        ws.column_dimensions[col[0].column_letter].width = max(10, min(22, len(str(col[0].value)) + 2))
    if nombre.startswith("Matriz"):                     # diagonal (aciertos) en verde y negrita
        for k in range(2, ws.max_row + 1):
            c = ws.cell(row=k, column=k)
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="C6EFCE")
            ws.cell(row=k, column=1).font = Font(bold=True)
    ws.freeze_panes = "B2" if nombre.startswith("Matriz") else "A2"
wb.save(SALIDA)

# ---- Mostrar en pantalla
print("=" * 78 + "\n RESUMEN COMPARATIVO (conjunto de prueba)\n" + "=" * 78)
for fila in resumen[1:]:
    print(f"{fila[0]:7s} accuracy {fila[1]:5.1f} ± {fila[2]:4.1f}%  |  F1 {fila[3]:5.1f} ± {fila[4]:4.1f}%  |  "
          f"{fila[5]:,} parámetros  |  {fila[6]} KB  |  {fila[7]} ms/seña")
print(f"Semillas: {', '.join(semillas)}")
print(f"\nGuardado: {SALIDA}  (7 hojas, incluidas las matrices de confusión)")